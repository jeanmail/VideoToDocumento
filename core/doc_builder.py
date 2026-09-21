"""
Construtor de documentos DOCX e PDF formatados para ingestão ideal no NotebookLM.
Suporta documento consolidado com a transcrição completa intercalada com as imagens aprovadas.
"""

import os
from typing import List, Optional
from PIL import Image

from core.video_processor import ExtractedFrame
from core.srt_parser import SubtitleItem

# DOCX
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

# PDF (ReportLab)
from reportlab.lib.pagesizes import letter
from reportlab.lib.units import inch
from reportlab.lib import colors
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Image as RLImage, Table, TableStyle, KeepTogether
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle


class DocumentBuilder:
    def __init__(self, title: str = "Guia de Treinamento - Passo a Passo", subtitle: str = "Documentação integral gerada automaticamente pelo VideoToDocument"):
        self.title = title
        self.subtitle = subtitle

    def _prepare_timeline(
        self,
        approved_frames: List[ExtractedFrame],
        all_subtitles: Optional[List[SubtitleItem]] = None
    ) -> List[dict]:
        """
        Organiza a linha do tempo cronológica intercalando transcrição completa e imagens aprovadas.
        Retorna lista de eventos: {'type': 'step_with_image', 'frame': frame, 'sub': sub}
        ou {'type': 'narration_only', 'sub': sub}
        """
        if not all_subtitles:
            # Se não houver transcrição completa, exibe apenas os frames aprovados
            return [{'type': 'step_with_image', 'frame': f, 'sub': None} for f in approved_frames]

        timeline = []
        pending_frames = sorted(approved_frames, key=lambda x: x.timestamp_seconds)
        sorted_subs = sorted(all_subtitles, key=lambda x: x.start_seconds)

        used_frames = set()

        for sub in sorted_subs:
            # Verifica se há frames anteriores ao início deste subtítulo
            while pending_frames and pending_frames[0].timestamp_seconds < sub.start_seconds:
                f = pending_frames.pop(0)
                used_frames.add(f.id)
                timeline.append({'type': 'step_with_image', 'frame': f, 'sub': None})

            # Verifica se há frame que incide dentro do intervalo deste subtítulo
            matched_frame = None
            if pending_frames and pending_frames[0].timestamp_seconds <= (sub.end_seconds + 0.5):
                matched_frame = pending_frames.pop(0)
                used_frames.add(matched_frame.id)

            if matched_frame:
                timeline.append({'type': 'step_with_image', 'frame': matched_frame, 'sub': sub})
            else:
                timeline.append({'type': 'narration_only', 'sub': sub, 'frame': None})

        # Adiciona qualquer frame restante
        while pending_frames:
            f = pending_frames.pop(0)
            if f.id not in used_frames:
                timeline.append({'type': 'step_with_image', 'frame': f, 'sub': None})

        return timeline

    def build_docx(
        self,
        approved_frames: List[ExtractedFrame],
        output_path: str,
        all_subtitles: Optional[List[SubtitleItem]] = None
    ) -> str:
        """
        Gera o documento Word (.docx) com a transcrição completa intercalada com as imagens aprovadas.
        """
        os.makedirs(os.path.dirname(os.path.abspath(output_path)), exist_ok=True)
        doc = docx.Document()

        # Configuração de Margens (1 polegada)
        for section in doc.sections:
            section.top_margin = Inches(1.0)
            section.bottom_margin = Inches(1.0)
            section.left_margin = Inches(1.0)
            section.right_margin = Inches(1.0)

        # Título Principal
        title_p = doc.add_heading(level=0)
        title_run = title_p.add_run(self.title)
        title_run.font.name = 'Calibri'
        title_run.font.size = Pt(24)
        title_run.font.bold = True
        title_run.font.color.rgb = RGBColor(24, 76, 120)

        # Subtítulo
        sub_p = doc.add_paragraph()
        sub_run = sub_p.add_run(self.subtitle)
        sub_run.font.name = 'Calibri'
        sub_run.font.size = Pt(11)
        sub_run.font.italic = True
        sub_run.font.color.rgb = RGBColor(100, 100, 100)

        doc.add_paragraph().paragraph_format.space_after = Pt(10)

        timeline = self._prepare_timeline(approved_frames, all_subtitles)
        step_counter = 0

        for item in timeline:
            if item['type'] == 'step_with_image':
                frame = item['frame']
                step_counter += 1
                step_name = frame.step_title if frame.step_title else f"Passo {step_counter}"
                
                # Cabeçalho semântico para NotebookLM
                h_text = f"[TIMESTAMP: {frame.timestamp_str}] - {step_name}"
                heading = doc.add_heading(level=2)
                h_run = heading.add_run(h_text)
                h_run.font.name = 'Calibri'
                h_run.font.size = Pt(13.5)
                h_run.font.bold = True
                h_run.font.color.rgb = RGBColor(30, 60, 90)
                heading.paragraph_format.space_before = Pt(14)
                heading.paragraph_format.space_after = Pt(6)

                # Imagem em alta resolução
                if os.path.exists(frame.image_path):
                    try:
                        p_img = doc.add_paragraph()
                        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
                        p_img.paragraph_format.space_after = Pt(6)
                        doc.add_picture(frame.image_path, width=Inches(5.5))
                    except Exception as e:
                        doc.add_paragraph(f"[Imagem: {e}]")

                # Transcrição associada ao print
                text_to_show = frame.subtitle_text or (item['sub'].text if item.get('sub') else "")
                if text_to_show:
                    p_sub = doc.add_paragraph()
                    p_sub.paragraph_format.left_indent = Inches(0.25)
                    p_sub.paragraph_format.space_after = Pt(4)
                    lbl = p_sub.add_run("Instrução / Fala:\n")
                    lbl.font.bold = True
                    lbl.font.size = Pt(10)
                    lbl.font.color.rgb = RGBColor(40, 40, 40)

                    txt = p_sub.add_run(text_to_show)
                    txt.font.size = Pt(10)
                    txt.font.color.rgb = RGBColor(30, 30, 30)

                # OCR (se houver)
                if frame.ocr_text:
                    p_ocr = doc.add_paragraph()
                    p_ocr.paragraph_format.left_indent = Inches(0.25)
                    p_ocr.paragraph_format.space_after = Pt(8)
                    ocr_lbl = p_ocr.add_run("Elementos da Interface (OCR): ")
                    ocr_lbl.font.bold = True
                    ocr_lbl.font.size = Pt(9)
                    ocr_lbl.font.color.rgb = RGBColor(90, 90, 90)

                    ocr_txt = p_ocr.add_run(frame.ocr_text)
                    ocr_txt.font.size = Pt(9)
                    ocr_txt.font.italic = True
                    ocr_txt.font.color.rgb = RGBColor(90, 90, 90)

            elif item['type'] == 'narration_only':
                sub = item['sub']
                if sub and sub.text:
                    p_narr = doc.add_paragraph()
                    p_narr.paragraph_format.left_indent = Inches(0.25)
                    p_narr.paragraph_format.space_after = Pt(4)
                    
                    time_lbl = p_narr.add_run(f"[{sub.start_time_str}] ")
                    time_lbl.font.bold = True
                    time_lbl.font.size = Pt(9.5)
                    time_lbl.font.color.rgb = RGBColor(70, 100, 130)

                    txt = p_narr.add_run(sub.text)
                    txt.font.size = Pt(9.5)
                    txt.font.color.rgb = RGBColor(45, 45, 45)

        doc.save(output_path)
        return output_path

    def build_pdf(
        self,
        approved_frames: List[ExtractedFrame],
        output_path: str,
        all_subtitles: Optional[List[SubtitleItem]] = None
    ) -> str:
        """
        Gera o documento no formato PDF com ReportLab com a transcrição completa intercalada com as imagens.
        """
        os.makedirs(os.path.dirname(os.path.abspath(output_path)), exist_ok=True)
        pdf_doc = SimpleDocTemplate(
            output_path,
            pagesize=letter,
            rightMargin=54,
            leftMargin=54,
            topMargin=54,
            bottomMargin=54
        )

        styles = getSampleStyleSheet()

        title_style = ParagraphStyle(
            'CustomTitle',
            parent=styles['Heading1'],
            fontName='Helvetica-Bold',
            fontSize=22,
            leading=26,
            textColor=colors.HexColor("#184C78"),
            spaceAfter=6
        )

        sub_style = ParagraphStyle(
            'CustomSubtitle',
            parent=styles['Normal'],
            fontName='Helvetica-Oblique',
            fontSize=10,
            leading=14,
            textColor=colors.HexColor("#646464"),
            spaceAfter=16
        )

        step_header_style = ParagraphStyle(
            'StepHeader',
            parent=styles['Heading2'],
            fontName='Helvetica-Bold',
            fontSize=12.5,
            leading=16,
            textColor=colors.HexColor("#1E3C5A"),
            spaceBefore=12,
            spaceAfter=6
        )

        body_style = ParagraphStyle(
            'StepBody',
            parent=styles['Normal'],
            fontName='Helvetica',
            fontSize=9.5,
            leading=13.5,
            textColor=colors.HexColor("#222222"),
            spaceAfter=6
        )

        narration_style = ParagraphStyle(
            'NarrationBody',
            parent=styles['Normal'],
            fontName='Helvetica',
            fontSize=9,
            leading=13,
            textColor=colors.HexColor("#333333"),
            leftIndent=14,
            spaceAfter=4
        )

        ocr_style = ParagraphStyle(
            'StepOCR',
            parent=styles['Normal'],
            fontName='Helvetica-Oblique',
            fontSize=8.5,
            leading=11.5,
            textColor=colors.HexColor("#555555"),
            spaceAfter=8
        )

        story = []
        story.append(Paragraph(self.title, title_style))
        story.append(Paragraph(self.subtitle, sub_style))

        max_img_width = 480
        timeline = self._prepare_timeline(approved_frames, all_subtitles)
        step_counter = 0

        for item in timeline:
            if item['type'] == 'step_with_image':
                frame = item['frame']
                step_counter += 1
                step_name = frame.step_title if frame.step_title else f"Passo {step_counter}"
                header_text = f"<b>[TIMESTAMP: {frame.timestamp_str}]</b> - {step_name}"

                step_elements = []
                step_elements.append(Paragraph(header_text, step_header_style))

                if os.path.exists(frame.image_path):
                    try:
                        with Image.open(frame.image_path) as img:
                            w, h = img.size
                            ratio = h / float(w)
                            target_w = min(w, max_img_width)
                            target_h = target_w * ratio

                        rl_img = RLImage(frame.image_path, width=target_w, height=target_h)
                        rl_img.hAlign = 'CENTER'
                        step_elements.append(rl_img)
                        step_elements.append(Spacer(1, 6))
                    except Exception:
                        pass

                text_to_show = frame.subtitle_text or (item['sub'].text if item.get('sub') else "")
                if text_to_show:
                    text_content = f"<b>Instrução / Fala:</b> {text_to_show}"
                    step_elements.append(Paragraph(text_content, body_style))

                if frame.ocr_text:
                    ocr_content = f"<b>Elementos Visuais (OCR):</b> {frame.ocr_text}"
                    step_elements.append(Paragraph(ocr_content, ocr_style))

                step_elements.append(Spacer(1, 8))
                story.append(KeepTogether(step_elements))

            elif item['type'] == 'narration_only':
                sub = item['sub']
                if sub and sub.text:
                    narr_text = f"<b><font color='#466482'>[{sub.start_time_str}]</font></b> {sub.text}"
                    story.append(Paragraph(narr_text, narration_style))

        pdf_doc.build(story)
        return output_path
